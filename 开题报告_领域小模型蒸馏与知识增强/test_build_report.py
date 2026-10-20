#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""build_report.py 的自检用例。

只覆盖解析与格式的关键不变量，不做完整渲染比对（docx 二进制含时间戳，不稳定）。

运行：
    ../.venv-docx/bin/python -m unittest -v test_build_report.py
"""

import sys
import unittest
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import build_report as br  # noqa: E402

from docx import Document  # noqa: E402
from docx.shared import Emu  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402


class TestParser(unittest.TestCase):
    """解析器不变量。"""

    def test_heading_levels(self):
        md = "# 题\n\n## 章\n\n### 节\n\n#### 小节\n"
        blocks, _ = br.parse_markdown(md)
        heads = [p for k, p in blocks if k == "heading"]
        self.assertEqual([h[0] for h in heads], [1, 2, 3, 4])
        self.assertEqual(heads[1][1], "章")

    def test_table_ragged_row_is_padded_not_dropped(self):
        """列数不齐的行必须保留（补齐），且发出警告。"""
        md = "| A | B | C |\n|---|---|---|\n| 1 | 2 | 3 |\n| 1 | 2 |\n| 4 | 5 | 6 |\n"
        blocks, warnings = br.parse_markdown(md)
        tables = [p for k, p in blocks if k == "table"]
        self.assertEqual(len(tables), 1)
        header, rows = tables[0]
        self.assertEqual(len(rows), 3, "三行都应保留")
        self.assertEqual(rows[1], ["1", "2", ""], "缺列应补空")
        self.assertTrue(warnings, "应产生警告")
        self.assertIn("列数", warnings[0])

    def test_table_extra_columns_truncated_with_warning(self):
        md = "| A | B |\n|---|---|\n| 1 | 2 | 3 |\n"
        blocks, warnings = br.parse_markdown(md)
        _, rows = [p for k, p in blocks if k == "table"][0]
        self.assertEqual(rows[0], ["1", "2"])
        self.assertTrue(warnings)

    def test_unknown_placeholder_warns(self):
        md = "正文\n\n{{FIG99}}\n"
        blocks, warnings = br.parse_markdown(md)
        self.assertTrue(any("未识别" in w for w in warnings))
        # 仍作为段落输出，不丢内容
        self.assertTrue(any(k == "para" and "{{FIG99}}" in p for k, p in blocks))

    def test_known_placeholder_no_warning(self):
        md = "{{FIG41}}\n"
        blocks, warnings = br.parse_markdown(md)
        self.assertEqual([k for k, _ in blocks], ["figure"])
        self.assertFalse(warnings)

    def test_table_separator_only_forms_table(self):
        """没有分隔行不应识别为表格。"""
        md = "| A | B |\n| 1 | 2 |\n"
        blocks, _ = br.parse_markdown(md)
        self.assertFalse(any(k == "table" for k, _ in blocks))


class TestSplit(unittest.TestCase):
    """正文与参考文献拆分。"""

    def test_title_and_refs(self):
        md = "# 题目X\n\n## 章\n\n正文\n\n## 参考文献\n\n[1] A.\n[2] B.\n"
        title, body, refs = br.split_markdown(md)
        self.assertEqual(title, "题目X")
        self.assertEqual(len(refs), 2)
        self.assertNotIn("参考文献", body)

    def test_no_refs_section(self):
        md = "# 题目\n\n## 章\n\n正文\n"
        title, body, refs = br.split_markdown(md)
        self.assertEqual(title, "题目")
        self.assertEqual(refs, [])

    def test_missing_title_uses_placeholder(self):
        title, _, _ = br.split_markdown("## 章\n正文\n")
        self.assertIn("未设置题目", title)


class TestBuildOutput(unittest.TestCase):
    """渲染产物的格式不变量。"""

    @classmethod
    def setUpClass(cls):
        import tempfile
        cls.tmp = tempfile.TemporaryDirectory()
        cls.out = Path(cls.tmp.name) / "t.docx"
        md = HERE / "开题报告_正文.md"
        if not md.exists():
            raise unittest.SkipTest("缺少 开题报告_正文.md")
        cls.n_refs, cls.n_heads, cls.warnings = br.build(md, cls.out)
        cls.doc = Document(str(cls.out))

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_page_setup(self):
        s = self.doc.sections[0]
        self.assertAlmostEqual(Emu(s.page_width).cm, 21.0, places=1)
        self.assertAlmostEqual(Emu(s.page_height).cm, 29.7, places=1)
        self.assertAlmostEqual(Emu(s.top_margin).cm, 2.54, places=2)
        self.assertAlmostEqual(Emu(s.left_margin).cm, 3.17, places=2)

    def test_heading_sizes(self):
        h1 = [p for p in self.doc.paragraphs if p.style.name == "Heading 1"]
        h2 = [p for p in self.doc.paragraphs if p.style.name == "Heading 2"]
        self.assertTrue(h1 and h2)
        self.assertEqual(Emu(h1[0].runs[0].font.size).pt, 15.0)
        self.assertEqual(Emu(h2[0].runs[0].font.size).pt, 14.0)

    def test_heading_color_not_blue(self):
        for p in self.doc.paragraphs:
            if not p.style.name.startswith("Heading") or not p.runs:
                continue
            col = p.runs[0].font.color
            if col is not None and col.rgb is not None:
                self.assertIn(str(col.rgb), ("000000", "FFFFFF"),
                              f"标题出现非黑色 {col.rgb}")

    def test_body_style(self):
        n = self.doc.styles["Normal"]
        self.assertEqual(Emu(n.font.size).pt, 12.0)
        self.assertAlmostEqual(Emu(n.paragraph_format.first_line_indent).cm, 0.847, places=2)
        self.assertEqual(n.paragraph_format.line_spacing, 1.3)

    def test_reference_numbering_continuous(self):
        import re
        nums = []
        for p in self.doc.paragraphs:
            m = re.match(r"^\[(\d+)\]", p.text.strip())
            if m:
                nums.append(int(m.group(1)))
        self.assertEqual(nums, list(range(1, len(nums) + 1)), "参考文献编号须连续")
        self.assertEqual(len(nums), self.n_refs)

    def test_citations_superscript(self):
        import re
        found = False
        for p in self.doc.paragraphs:
            for r in p.runs:
                if re.fullmatch(r"\[\d+\](?:\[\d+\])*", r.text) and r.font.superscript:
                    found = True
                    break
            if found:
                break
        self.assertTrue(found, "正文引用应渲染为上标")

    def test_table_has_no_empty_leading_run(self):
        """表格单元格不应因 cell.text=\"\" 留下空 run。"""
        self.assertTrue(self.doc.tables)
        cell = self.doc.tables[0].rows[0].cells[0]
        para = cell.paragraphs[0]
        self.assertEqual(len(para.runs), 1, f"应只有 1 个 run，实际 {len(para.runs)}")
        self.assertTrue(para.text.strip())

    def test_cover_info_blank(self):
        """封面个人信息必须留空（下划线占位）。"""
        head = [p.text for p in self.doc.paragraphs[:20] if p.text.strip()]
        joined = "\n".join(head)
        self.assertIn("北京科技大学", joined.replace(" ", ""))
        self.assertIn("选题报告及文献综述", joined)
        self.assertIn("＿", joined, "信息栏应有下划线占位")

    def test_document_ends_with_content_not_empty_para(self):
        body = self.doc.element.body
        kids = list(body)
        # 最后一个非 sectPr 元素不应是纯空段落
        tail = [k for k in kids if k.tag != qn("w:sectPr")]
        self.assertTrue(tail)
        last = tail[-1]
        if last.tag == qn("w:p"):
            txt = "".join(t.text or "" for t in last.findall(".//" + qn("w:t")))
            self.assertTrue(txt.strip(), "末尾不应是空段落")


class TestCLIContract(unittest.TestCase):
    """命令行契约：默认输出名按模式区分。"""

    def test_default_names_differ(self):
        self.assertNotEqual(br.DEFAULT_OUT_TOC, br.DEFAULT_OUT_STATIC)


class TestGlobalState(unittest.TestCase):
    """parse_markdown 之间不得共享警告状态（无模块级全局）。"""

    def test_warnings_do_not_accumulate(self):
        md = "| A |\n|---|\n| 1 | 2 |\n"
        _, w1 = br.parse_markdown(md)
        _, w2 = br.parse_markdown(md)
        self.assertEqual(len(w1), 1)
        self.assertEqual(len(w2), 1, "第二次调用不应累积上一次的警告")

    def test_no_module_level_warnings_attr(self):
        self.assertFalse(hasattr(br, "WARNINGS"),
                         "不应再有模块级 WARNINGS 全局")

    def test_inline_unknown_placeholder_warns(self):
        """行内占位符也要告警（原实现只识别整行）。"""
        md = "见图 {{FIG99}} 所示。\n"
        _, warnings = br.parse_markdown(md)
        self.assertTrue(any("未识别" in w for w in warnings),
                        "行内占位符应告警")


class TestCLI(unittest.TestCase):
    """main() 的命令行契约（此前完全没有测试）。"""

    def setUp(self):
        import tempfile
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.md = self.dir / "in.md"
        self.md.write_text("# 题\n\n## 章\n\n正文\n\n## 参考文献\n\n[1] A.\n",
                           encoding="utf-8")

    def tearDown(self):
        self.tmp.cleanup()

    def _run(self, *args):
        """调用 main，返回 (退出码, stdout, stderr)。"""
        import contextlib
        import io
        out, err = io.StringIO(), io.StringIO()
        code = 0
        with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
            try:
                br.main(["build_report.py", *args])
            except SystemExit as e:
                code = e.code if isinstance(e.code, int) else 1
        return code, out.getvalue(), err.getvalue()

    def test_help_exits_zero(self):
        code, out, _ = self._run("--help")
        self.assertEqual(code, 0)
        self.assertIn("用法", out)
        self.assertIn("--fig-caption", out)

    def test_unknown_flag_exits_two(self):
        code, _, err = self._run("--bogus")
        self.assertEqual(code, 2)
        self.assertIn("未知参数", err)

    def test_too_many_positionals_exits_two(self):
        code, _, err = self._run("a.md", "b.docx", "c.docx")
        self.assertEqual(code, 2)
        self.assertIn("参数过多", err)

    def test_missing_input_exits_one(self):
        code, _, err = self._run(str(self.dir / "nope.md"), str(self.dir / "o.docx"))
        self.assertEqual(code, 1)
        self.assertIn("找不到输入文件", err)

    def test_refuses_overwrite_without_force(self):
        out = self.dir / "o.docx"
        code1, _, _ = self._run(str(self.md), str(out))
        self.assertEqual(code1, 0)
        self.assertTrue(out.exists())
        code2, _, err = self._run(str(self.md), str(out))
        self.assertEqual(code2, 2, "已存在且无 --force 应拒绝")
        self.assertIn("已存在", err)

    def test_force_overwrites(self):
        out = self.dir / "o.docx"
        self._run(str(self.md), str(out))
        code, _, _ = self._run(str(self.md), str(out), "--force")
        self.assertEqual(code, 0)

    def test_refuses_same_input_output_path(self):
        """输入输出同路径会销毁输入 md，必须拒绝。"""
        before = self.md.read_text(encoding="utf-8")
        code, _, err = self._run(str(self.md), str(self.md), "--force")
        self.assertEqual(code, 2)
        self.assertIn("同一个文件", err)
        self.assertEqual(self.md.read_text(encoding="utf-8"), before,
                         "输入文件内容不得被改动")

    def test_missing_output_dir_exits_one(self):
        code, _, err = self._run(str(self.md), str(self.dir / "no" / "such" / "o.docx"))
        self.assertEqual(code, 1)
        self.assertIn("输出目录不存在", err)

    def test_default_output_names_differ_by_mode(self):
        code1, out1, _ = self._run(str(self.md))
        code2, out2, _ = self._run(str(self.md), "--static-toc")
        # 两个模式默认都写本目录，可能与本机已有文件冲突，这里只校验名字约定不同
        self.assertNotEqual(br.DEFAULT_OUT_TOC, br.DEFAULT_OUT_STATIC)

    def test_fig_caption_option_applied(self):
        """--fig-caption 应被接受并写入文档图题。"""
        out = self.dir / "o.docx"
        code, _, _ = self._run(str(self.md), str(out), "--fig-caption", "图 9-9　自定义")
        self.assertEqual(code, 0)
        self.assertTrue(out.exists())
        # 无 {{FIG}} 占位时不会插图，这里只验证参数不报错

    def test_fig_caption_missing_value_exits_two(self):
        code, _, err = self._run(str(self.md), "--fig-caption")
        self.assertEqual(code, 2)
        self.assertIn("缺少取值", err)


class TestFormatCoverage(unittest.TestCase):
    """此前未覆盖的格式断言：目录域、页码域、图题、表格样式。"""

    @classmethod
    def setUpClass(cls):
        import tempfile
        cls.tmp = tempfile.TemporaryDirectory()
        cls.out = Path(cls.tmp.name) / "t.docx"
        md = HERE / "开题报告_正文.md"
        if not md.exists():
            raise unittest.SkipTest("缺少 开题报告_正文.md")
        fig = None
        for cand in ("技术路线.png", "图4-1_技术路线图.png", "图3-1_技术路线图.png"):
            if (HERE / cand).exists():
                fig = HERE / cand
                break
        cls.n_refs, cls.n_heads, cls.warnings = br.build(md, cls.out, fig_path=fig)
        cls.doc = Document(str(cls.out))

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_toc_field_present_in_toc_mode(self):
        joined = "\n".join(p._p.xml for p in self.doc.paragraphs)
        self.assertIn("TOC", joined, "TOC 域应存在")

    def test_page_number_field_in_footer(self):
        footer = self.doc.sections[0].footer
        xml = "\n".join(p._p.xml for p in footer.paragraphs)
        self.assertIn("PAGE", xml, "页脚应有 PAGE 域")

    def test_figure_caption_present(self):
        caps = [p.text.strip() for p in self.doc.paragraphs
                if p.text.strip().startswith("图 ")]
        self.assertTrue(caps, "应有图题")

    def test_table_style_is_grid(self):
        self.assertTrue(self.doc.tables)
        self.assertEqual(self.doc.tables[0].style.name, "Table Grid")


if __name__ == "__main__":
    unittest.main(verbosity=2)
