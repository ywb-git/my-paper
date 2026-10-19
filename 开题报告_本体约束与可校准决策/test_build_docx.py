import importlib.util
from pathlib import Path
import tempfile
import unittest

from docx import Document


MODULE_PATH = Path(__file__).with_name("build_docx.py")
spec = importlib.util.spec_from_file_location("report_builder", MODULE_PATH)
builder = importlib.util.module_from_spec(spec)
spec.loader.exec_module(builder)


class BuildDocxTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.original_md, self.original_out = builder.MD, builder.OUT
        self.addCleanup(self.restore_paths)
        builder.MD = Path(self.temp.name) / "source.md"
        builder.OUT = Path(self.temp.name) / "report.docx"

    def restore_paths(self):
        builder.MD, builder.OUT = self.original_md, self.original_out

    def build(self, source):
        builder.MD.write_text(source, encoding="utf-8")
        builder.build()
        return Document(builder.OUT)

    def test_title_and_contents_follow_source(self):
        doc = self.build("# 修订后的题目\n\n## 新章\n\n### 新小节\n\n正文。\n")
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("论文题目：修订后的题目", text)
        self.assertEqual(text.count("新章"), 2)
        self.assertEqual(text.count("新小节"), 2)
        self.assertNotIn("3.1.0", text)

    def test_markdown_table_is_a_word_table(self):
        doc = self.build("# 表格测试\n\n## 比较\n\n| 方案 | 结果 |\n|---|---|\n| 全局 | 基线 |\n| 分组 | 待检验 |\n")
        self.assertEqual(len(doc.tables), 1)
        self.assertEqual([[c.text for c in row.cells] for row in doc.tables[0].rows],
                         [["方案", "结果"], ["全局", "基线"], ["分组", "待检验"]])
        self.assertFalse(any("|---" in p.text for p in doc.paragraphs))


if __name__ == "__main__":
    unittest.main()
